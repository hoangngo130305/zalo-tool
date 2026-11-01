import hmac
import io
import json
from datetime import datetime
from functools import wraps

from django import forms
from django.conf import settings
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse, JsonResponse
from django.db import transaction
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods

from . import services as sv
from . import worker
from .auth import con_bi_khoa
from .models import (CaiDat, CauHinhNhom, KetQua, KhongLienHe, LichSuGui, LichSuGuiNhom, NguoiDung, NguoiNhan, NhatKy,
                     chuan_hoa_sdt, q_thanh_cong)


# ---------- trang ----------
class DangNhapForm(AuthenticationForm):
    username = forms.EmailField(label="Email", widget=forms.EmailInput(attrs={"autofocus": True, "autocomplete": "email"}))
    error_messages = {
        "invalid_login": "Sai email hoặc mật khẩu.",
        "inactive": "Tài khoản này đã bị khoá.",
    }

    def clean(self):
        email = self.cleaned_data.get("username") or ""
        phut = con_bi_khoa(self.request, email)
        if phut:
            raise forms.ValidationError(f"Sai quá nhiều lần, thử lại sau {phut} phút.")
        try:
            return super().clean()
        except PermissionDenied:
            raise forms.ValidationError(f"Sai quá nhiều lần, thử lại sau {con_bi_khoa(self.request, email) or 5} phút.") from None


class DangNhapView(auth_views.LoginView):
    template_name = "registration/login.html"
    authentication_form = DangNhapForm
    redirect_authenticated_user = True

    def form_valid(self, form):
        res = super().form_valid(form)
        u = self.request.user
        sv.ghi_log(f"{u} ({u.get_vai_tro_display()}) đã đăng nhập.", nguoi_dung=u)
        return res


@login_required
def trang_chu(request):
    """Quản trị / chỉ xem làm việc trong trang quản trị; nhân viên dùng trang gửi tin riêng."""
    u = request.user
    if u.is_staff:
        return redirect("admin:index")
    me = {"id": u.id, "name": str(u), "email": u.email, "role": u.vai_tro, "roleLabel": u.get_vai_tro_display(),
          "perms": sorted(u.quyen)}
    return render(request, "guitin/nhan_vien.html", {"me": me})


# ---------- khung API ----------
def api(quyen="view", methods=("GET", "POST")):
    """Bắt buộc đăng nhập + có quyền; trả lỗi dạng JSON để trang hiển thị."""
    def deco(fn):
        @wraps(fn)
        @require_http_methods(list(methods))
        def wrapper(request, *args, **kwargs):
            u = request.user
            if not u.is_authenticated:
                return JsonResponse({"error": "Phiên đăng nhập đã hết, hãy đăng nhập lại"}, status=401)
            if not u.co_quyen(quyen):
                return JsonResponse({"error": "Tài khoản của bạn không có quyền làm việc này"}, status=403)
            try:
                if request.method == "POST" and request.content_type == "application/json":
                    request.json = json.loads(request.body or b"{}")
                else:
                    request.json = {}
                out = fn(request, *args, **kwargs)
                return out if isinstance(out, HttpResponse) else JsonResponse(out)
            except (ValueError, worker.WorkerLoi) as e:
                return JsonResponse({"error": str(e)}, status=400)
        return wrapper
    return deco


def trang_thai_cho(user):
    """Trạng thái Zalo của chính người đang đăng nhập (mỗi người một tài khoản Zalo riêng)."""
    return worker.trang_thai(user.id)


def staff_cho(user):
    return sv.staff_dict(sv.nguoi_duoc_giao()) if user.co_quyen("all") else []


def can_khong_dang_gui(user):
    if worker.trang_thai(user.id).get("running"):
        raise ValueError("Đang gửi, hãy dừng hoặc chờ xong rồi làm tiếp")


def han_muc_cho(user):
    """Hạn mức hôm nay của Zalo người này (hiện trên trang gửi tin)."""
    from django.utils.timezone import localtime
    cfg = CaiDat.lay()
    tam = user.tam_dung_den if user.tam_dung_den and user.tam_dung_den > timezone.now() else None
    return {"daGui": sv.da_gui_hom_nay(user), "toiDa": sv.han_muc_ngay(user, cfg),
            "tamDungDen": localtime(tam).strftime("%H:%M %d/%m") if tam else None}


def nhat_ky_cho(user):
    """Nhân viên chỉ thấy nhật ký của mình (thao tác của mình và các đợt gửi mình bấm)."""
    qs = NhatKy.objects.all()
    return qs if user.co_quyen("all") else qs.filter(nguoi_dung=user)


# ---------- API ----------
@api("view", ["GET"])
def api_state(request):
    logs = list(nhat_ky_cho(request.user).order_by("-id")[:300])
    return {
        **trang_thai_cho(request.user),
        "config": CaiDat.lay().to_dict(),
        "tinNhanCuaToi": request.user.tin_nhan_mac_dinh,
        "loiMoiCuaToi": request.user.loi_moi_ket_ban,
        "anhCuaToi": request.user.anh_tin_nhan,
        "contacts": sv.danh_sach(request.user),
        "staff": staff_cho(request.user),
        "logs": [sv.log_dict(n) for n in reversed(logs)],
        "hanMuc": han_muc_cho(request.user),
        "version": sv.phien_ban(request.user),
    }


@api("view", ["GET"])
def api_poll(request):
    """Trang gọi mỗi 1,5 giây: trạng thái gửi, nhật ký mới, và phiên bản danh sách."""
    try:
        sau = int(request.GET.get("log") or 0)
    except ValueError:
        sau = 0
    logs = nhat_ky_cho(request.user).filter(id__gt=sau).order_by("id")[:200]
    return {**trang_thai_cho(request.user), "logs": [sv.log_dict(n) for n in logs], "hanMuc": han_muc_cho(request.user),
            "version": sv.phien_ban(request.user)}


@api("all", ["GET"])
def api_thong_ke(request):
    """Số liệu cho bảng điều khiển trong trang quản trị."""
    return {**sv.thong_ke(), "version": sv.phien_ban(request.user)}


@api("view", ["GET", "POST"])
def api_contacts(request):
    ket_qua_luu = None
    if request.method == "POST":
        if not request.user.co_quyen("list"):
            return JsonResponse({"error": "Tài khoản của bạn không có quyền sửa danh sách"}, status=403)
        # Quản trị sửa danh sách trong trang quản trị; đường này chỉ để nhân viên lưu danh sách của mình
        if request.user.co_quyen("all"):
            return JsonResponse({"error": "Quản trị sửa danh sách trong trang quản trị → Người nhận"}, status=403)
        can_khong_dang_gui(request.user)
        rows = request.json.get("contacts") or []
        if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
            raise ValueError("Dữ liệu không hợp lệ")
        ket_qua_luu = sv.luu_danh_sach(request.user, rows)
    return {"contacts": sv.danh_sach(request.user), "luu": ket_qua_luu, "staff": staff_cho(request.user), "config": CaiDat.lay().to_dict(),
            "tinNhanCuaToi": request.user.tin_nhan_mac_dinh, "loiMoiCuaToi": request.user.loi_moi_ket_ban,
            "anhCuaToi": request.user.anh_tin_nhan, "version": sv.phien_ban(request.user)}


@api("send", ["POST"])
def api_tin_nhan_cua_toi(request):
    """Nhân viên tự sửa tin nhắn mặc định gửi hàng loạt của mình (để trống = dùng nội dung của quản trị)."""
    u = request.user
    u.tin_nhan_mac_dinh = str(request.json.get("noiDung") or "").strip()[:5000]
    u.save(update_fields=["tin_nhan_mac_dinh"])
    sv.tang_phien_ban()
    return {"tinNhanCuaToi": u.tin_nhan_mac_dinh, "version": sv.phien_ban(request.user)}


@api("send", ["POST"])
def api_loi_moi_cua_toi(request):
    """Nhân viên tự sửa lời mời kết bạn của mình (để trống = dùng lời mời của quản trị)."""
    u = request.user
    u.loi_moi_ket_ban = str(request.json.get("loiMoi") or "").strip()[:3000]
    u.save(update_fields=["loi_moi_ket_ban"])
    sv.tang_phien_ban()
    return {"loiMoiCuaToi": u.loi_moi_ket_ban, "version": sv.phien_ban(request.user)}


@api("send", ["POST"])
def api_anh_cua_toi(request):
    """Nhân viên tải ảnh gửi kèm tin nhắn của mình (multipart, trường "anh", có thể nhiều file)."""
    can_khong_dang_gui(request.user)
    files = request.FILES.getlist("anh")
    if not files:
        raise ValueError("Chưa chọn ảnh")
    for f in files:
        sv.luu_anh_nv(request.user, f)
    sv.tang_phien_ban()
    return {"anhCuaToi": request.user.anh_tin_nhan, "version": sv.phien_ban(request.user)}


@api("send", ["POST"])
def api_anh_cua_toi_xoa(request):
    can_khong_dang_gui(request.user)
    sv.xoa_anh_nv(request.user, str(request.json.get("ten") or ""))
    sv.tang_phien_ban()
    return {"anhCuaToi": request.user.anh_tin_nhan, "version": sv.phien_ban(request.user)}


@login_required
def xem_anh_cua_toi(request, ten):
    from django.http import FileResponse, Http404
    if ten not in request.user.anh_tin_nhan:
        raise Http404
    f = sv.thu_muc_anh_nv(request.user) / ten
    if not f.is_file():
        raise Http404
    return FileResponse(open(f, "rb"))


@api("list", ["POST"])
def api_tu_choi(request):
    """Khách từ chối liên hệ: thêm SĐT vào danh sách không liên hệ (không gửi nữa, nhập lại cũng bị loại).
    Nhân viên chỉ thêm được số trong danh sách của mình. bo=true để bỏ khỏi danh sách (chỉ quản trị)."""
    u = request.user
    sdt = chuan_hoa_sdt(request.json.get("sdt"))
    if not sdt:
        raise ValueError("Số điện thoại không hợp lệ")
    if request.json.get("bo"):
        if not u.co_quyen("settings"):
            raise ValueError("Chỉ quản trị được bỏ số khỏi danh sách từ chối")
        KhongLienHe.objects.filter(sdt=sdt).delete()
        sv.ghi_log(f"{u} đã bỏ {sdt} khỏi danh sách khách từ chối liên hệ.", "warn", u)
    else:
        if not u.co_quyen("all") and not u.nguoi_nhan.filter(sdt=sdt).exists():
            raise ValueError("Số này không thuộc danh sách của bạn")
        ly_do = str(request.json.get("lyDo") or "").strip()[:255] or "Khách từ chối liên hệ"
        KhongLienHe.objects.get_or_create(sdt=sdt, defaults={"ly_do": ly_do, "nguoi_them": u})
        sv.ghi_log(f"{u} đánh dấu {sdt} là khách từ chối liên hệ ({ly_do}). Tool sẽ không gửi cho số này nữa.", "warn", u)
    sv.tang_phien_ban()
    return {"contacts": sv.danh_sach(u), "version": sv.phien_ban(u)}


@api("list", ["POST"])
def api_import(request):
    return {"rows": sv.doc_file(request.body)}


@api("view", ["POST"])
def api_preview(request):
    cfg = CaiDat.lay()
    chan = sv.chua_gui_duoc(request.user, cfg)
    if chan:
        raise ValueError(chan)
    items = sv.cho_gui(request.user, request.json.get("phamVi"), sv.so_gui_toi_da(request.user, cfg))
    so_anh = len(sv.anh_gui_kem(request.user, cfg)) if cfg.gui_tin_nhan else 0
    # Có nhiều mẫu thì xem trước một mẫu ngẫu nhiên (lúc gửi thật mỗi người cũng nhận ngẫu nhiên 1 mẫu)
    out = []
    for r in items:
        loi_moi, tin = sv.dien(sv.loi_moi_gui(r, cfg), r.ten), sv.noi_dung_gui(r, cfg)
        out.append({
            "sdt": r.sdt, "ten": r.ten,
            "loiMoi": sv.xao(loi_moi) if cfg.gui_ket_ban else None, "soMauLoiMoi": len(sv.cac_mau(loi_moi)),
            "tenGoiNho": sv.ten_goi_nho(cfg.mau_ten_goi_nho, r.ten, r.sdt, sv.du_an_cua(r, cfg)) if cfg.doi_ten_goi_nho else None,
            "tinNhan": sv.xao(tin) if cfg.gui_tin_nhan else None, "soMauTinNhan": len(sv.cac_mau(tin)),
            "soAnh": so_anh,
        })
    return {"items": out, "hanMuc": han_muc_cho(request.user)}


@api("send", ["POST"])
def api_reset(request):
    """Gửi lại 1 số (nhân viên chỉ với số của mình), hoặc xoá toàn bộ lịch sử (chỉ quản trị)."""
    can_khong_dang_gui(request.user)
    u = request.user
    sdt = request.json.get("sdt")
    if not sdt:
        if not u.co_quyen("settings"):
            raise ValueError("Chỉ quản trị viên được xoá toàn bộ lịch sử gửi")
        KetQua.objects.all().delete()
        sv.ghi_log(f"{u} đã xoá toàn bộ lịch sử gửi.", "warn", u)
    else:
        phone = chuan_hoa_sdt(sdt)
        if not u.co_quyen("all") and not u.nguoi_nhan.filter(sdt=phone).exists():
            raise ValueError("Số này không thuộc danh sách của bạn")
        KetQua.objects.filter(sdt=phone).delete()
        sv.ghi_log(f"{u} đánh dấu {phone} để gửi lại.", "warn", u)
    sv.tang_phien_ban()
    return {"contacts": sv.danh_sach(u), "version": sv.phien_ban(request.user)}


# ---------- Zalo của từng người ----------
def _chu_tai_khoan(request):
    """Tài khoản Zalo được thao tác: mặc định của chính mình; quản trị thao tác hộ được người khác."""
    uid = request.json.get("uid")
    if uid and int(uid) != request.user.id:
        if not request.user.co_quyen("settings"):
            raise ValueError("Bạn chỉ thao tác được với Zalo của chính mình")
        chu = NguoiDung.objects.filter(id=int(uid)).first()
        if not chu:
            raise ValueError("Không tìm thấy tài khoản này")
        return chu
    return request.user


@api("send", ["POST"])
def api_zalo_login(request):
    worker.goi("POST", "/login", {"uid": request.user.id, "qr": bool(request.json.get("qr"))})
    return {"ok": True}


@api("send", ["POST"])
def api_zalo_cancel(request):
    worker.goi("POST", "/login/cancel", {"uid": request.user.id})
    return {"ok": True}


@api("send", ["POST"])
def api_zalo_logout(request):
    chu = _chu_tai_khoan(request)
    worker.goi("POST", "/logout", {"uid": chu.id})
    ai = "Zalo của mình" if chu == request.user else f"Zalo của {chu}"
    sv.ghi_log(f"{request.user} đã ngắt kết nối {ai}.", nguoi_dung=request.user)
    return {"ok": True}


@api("all", ["GET"])
def api_zalo_tat_ca(request):
    """Zalo của từng nhân viên: đã kết nối chưa, đang gửi tới đâu (cho bảng điều khiển quản trị)."""
    st = worker.tat_ca()
    out = []
    for u in sv.nguoi_duoc_giao():
        s = (st or {}).get(u.id, {})
        out.append({"id": u.id, "name": str(u), "role": u.vai_tro, "login": s.get("login", "out"),
                    "account": s.get("account"), "running": bool(s.get("running")), "progress": s.get("progress"),
                    "current": s.get("current"), "stopRequested": bool(s.get("stopRequested"))})
    return {"worker": st is not None, "nhanVien": out}


@api("send", ["POST"])
def api_start(request):
    """Gửi bằng Zalo của chính người bấm gửi."""
    u = request.user
    if worker.trang_thai(u.id).get("login") != "in":
        raise ValueError("Bạn chưa kết nối Zalo. Quét QR để kết nối Zalo của bạn trước khi gửi")
    cfg = CaiDat.lay()
    if not cfg.gui_ket_ban and not cfg.gui_tin_nhan:
        raise ValueError("Đang tắt cả kết bạn lẫn tin nhắn, hãy bật ít nhất một mục")
    chan = sv.chua_gui_duoc(u, cfg)
    if chan:
        raise ValueError(chan)
    batch = sv.cho_gui(u, request.json.get("phamVi"), sv.so_gui_toi_da(u, cfg))
    if not batch:
        raise ValueError("Không còn ai cần gửi trong phạm vi này")
    worker.goi("POST", "/run", {
        "uid": u.id,
        # Gửi kèm nội dung đã điền sẵn cho từng người (nội dung riêng hoặc mặc định)
        "batch": [{"sdt": r.sdt, "ten": r.ten, "noiDung": sv.noi_dung_gui(r, cfg), "duAn": sv.du_an_cua(r, cfg),
                   "loiMoi": sv.loi_moi_gui(r, cfg), "nhanVien": sv.nhan_vien_cua(r)} for r in batch],
        # Ảnh gửi kèm: ảnh riêng của người bấm gửi, chưa có thì ảnh của quản trị.
        # conLaiHomNay: worker tự dừng khi đủ hạn mức hôm nay (lớp chặn thứ hai, cả khi đợt gửi kéo qua nửa đêm)
        "cfg": {**cfg.to_dict(), "anh": sv.anh_gui_kem(u, cfg), "conLaiHomNay": sv.con_lai_hom_nay(u, cfg),
                "hanMucNgay": sv.han_muc_ngay(u, cfg)},
        "runBy": {"id": u.id, "name": str(u)},
    })
    return {"ok": True}


@api("send", ["POST"])
def api_stop(request):
    """Dừng đợt gửi của mình; quản trị dừng hộ được đợt gửi của nhân viên."""
    chu = _chu_tai_khoan(request)
    worker.goi("POST", "/stop", {"uid": chu.id})
    ai = "" if chu == request.user else f" đợt gửi của {chu}"
    sv.ghi_log(f"{request.user} bấm dừng{ai}.", "warn", request.user)
    return {"ok": True}


@api("view", ["GET"])
def api_report(request):
    """Báo cáo Excel. Nhân viên chỉ tải được kết quả các số của mình."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    qs = KetQua.objects.select_related("nguoi_gui").order_by("thoi_gian")
    if not request.user.co_quyen("all"):
        qs = qs.filter(sdt__in=request.user.nguoi_nhan.exclude(sdt__isnull=True).values("sdt"))
    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    head = ["SĐT", "Tên trong danh sách", "Tên Zalo", "Kết bạn", "Tin nhắn", "Ghi chú", "Thời gian", "Người bấm gửi"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FC4")
    n = 0
    for k in qs:
        sv.dong_excel(ws, [k.sdt, k.ten, k.ten_zalo, k.ket_ban, k.tin_nhan, k.ghi_chu,
                           timezone.localtime(k.thoi_gian).strftime("%H:%M:%S %d/%m/%Y"), str(k.nguoi_gui or "")])
        n += 1
    for col, w in zip("ABCDEFGH", (13, 24, 24, 12, 12, 45, 20, 20)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    sv.ghi_log(f"{request.user} đã xuất Excel kết quả gửi ({n} dòng).", nguoi_dung=request.user)
    res = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    res["Content-Disposition"] = f'attachment; filename="ket-qua-{timezone.localdate():%Y%m%d}.xlsx"'
    return res


@login_required
def file_mau(request):
    """File Excel mẫu để nhân viên điền: SĐT, Họ tên, Nội dung."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "DanhSach"
    ws.append(["SĐT", "Họ tên", "Nội dung", "Dự án"])
    ws.append(["0901234567", "Anh Nam", "Chào {ten}, em gửi anh thông tin chương trình tháng này ạ.", "Palm City"])
    ws.append(["0912345678", "Chị Lan", "", "Global City"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FC4")
    for r in (2, 3):
        ws.cell(r, 1).number_format = "@"   # giữ số 0 đầu
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws["F1"] = "Ghi chú"
    ws["F2"] = "{ten} sẽ được thay bằng họ tên. Để trống Nội dung thì gửi nội dung mặc định."
    ws["F3"] = "Dự án được ghi vào tên gợi nhớ khi kết bạn. Để trống thì dùng dự án mặc định của quản trị."
    ws["F1"].font = Font(bold=True)
    for col, w in zip("ABCDEF", (14, 22, 60, 18, 2, 60)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    res = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    res["Content-Disposition"] = 'attachment; filename="mau-danh-sach.xlsx"'
    return res


# ---------- lịch sử gửi ----------
def _bo_loc(request):
    g = request.GET
    return {k: (g.get(k) or "").strip() for k in ("q", "tt", "tu", "den", "nguoi")}


@api("view", ["GET"])
def api_lich_su(request):
    """Lịch sử gửi có lọc & phân trang. Nhân viên chỉ thấy lịch sử mình gửi."""
    qs = sv.loc_lich_su(request.user, **_bo_loc(request))
    try:
        trang = max(1, int(request.GET.get("trang") or 1))
    except ValueError:
        trang = 1
    co = 50
    tong = qs.count()
    ok = qs.filter(q_thanh_cong()).count()
    rows = qs[(trang - 1) * co: trang * co]
    return {"rows": [sv.lich_su_dict(h) for h in rows], "tong": tong, "thanhCong": ok, "loi": tong - ok,
            "trang": trang, "soTrang": max(1, (tong + co - 1) // co)}


@api("view", ["GET"])
def api_lich_su_excel(request):
    """Xuất Excel đúng theo bộ lọc đang chọn."""
    loc = _bo_loc(request)
    qs = sv.loc_lich_su(request.user, **loc)
    data = sv.excel_lich_su(qs)
    mo_ta = ", ".join(f"{k}={v}" for k, v in loc.items() if v) or "không lọc"
    sv.ghi_log(f"{request.user} đã xuất Excel lịch sử gửi ({qs.count()} dòng, {mo_ta}).", nguoi_dung=request.user)
    res = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    res["Content-Disposition"] = f'attachment; filename="lich-su-gui-{timezone.localdate():%Y%m%d}.xlsx"'
    return res


# ---------- worker báo kết quả về ----------
@csrf_exempt
@require_http_methods(["POST"])
def api_worker_event(request):
    got = request.headers.get("X-Worker-Token", "")
    if not hmac.compare_digest(got, settings.WORKER_TOKEN):
        return JsonResponse({"error": "Sai mã worker"}, status=403)
    b = json.loads(request.body or b"{}")
    if b.get("type") == "auto_stop":
        # Worker tự dừng vì lỗi liên tiếp: tạm dừng Zalo này một thời gian để tài khoản không bị phạt nặng hơn
        cfg = CaiDat.lay()
        u = NguoiDung.objects.filter(id=b.get("nguoiDungId")).first()
        if u and cfg.tam_dung_gio:
            from datetime import timedelta
            u.tam_dung_den = timezone.now() + timedelta(hours=cfg.tam_dung_gio)
            u.save(update_fields=["tam_dung_den"])
            sv.ghi_log(f"Zalo của {u} bị tạm dừng gửi {cfg.tam_dung_gio} giờ vì gặp {b.get('soLoi')} lỗi liên tiếp. "
                       f"Kiểm tra Zalo trên điện thoại; quản trị bỏ tạm dừng trong Tài khoản nếu muốn gửi lại sớm.", "error", u)
            sv.tang_phien_ban([u.id])
    elif b.get("type") == "log":
        NhatKy.objects.create(noi_dung=str(b.get("text", ""))[:5000], muc=b.get("level") or "info",
                              nguoi_dung=NguoiDung.objects.filter(id=b.get("nguoiDungId")).first())
    elif b.get("type") == "result":
        sdt = chuan_hoa_sdt(b.get("sdt"))
        r = b.get("ketQua") or {}
        if sdt:
            try:
                tg = datetime.fromisoformat(str(r.get("thoiGian")).replace("Z", "+00:00"))
            except ValueError:
                tg = timezone.now()
            du_lieu = {
                "ten": str(r.get("ten") or "")[:255], "ten_zalo": str(r.get("tenZalo") or "")[:255],
                "ket_ban": str(r.get("ketBan") or "-")[:20], "tin_nhan": str(r.get("tinNhan") or "-")[:20],
                "noi_dung": str(r.get("noiDung") or ""), "ghi_chu": str(r.get("ghiChu") or ""), "thoi_gian": tg,
                "tai_khoan_zalo": str(r.get("taiKhoanZalo") or "")[:255],
                "nguoi_gui": NguoiDung.objects.filter(id=b.get("nguoiGuiId")).first(),
            }
            # Worker gửi lại kết quả chưa được xác nhận (hàng đợi trên đĩa): mã đã có thì bỏ qua, không ghi 2 lần
            ma = str(b.get("ma") or "")[:40] or None
            if not (ma and LichSuGui.objects.filter(ma=ma).exists()):
                with transaction.atomic():
                    # Một dòng lịch sử (không ghi đè) + trạng thái mới nhất của số này
                    LichSuGui.objects.create(sdt=sdt, ma=ma, **du_lieu)
                    KetQua.objects.update_or_create(sdt=sdt, defaults=du_lieu)
            phu_trach = NguoiNhan.objects.filter(sdt=sdt).values_list("phu_trach_id", flat=True)
            sv.tang_phien_ban([b.get("nguoiGuiId"), *phu_trach])
    elif b.get("type") == "result_nhom":
        r = b.get("ketQua") or {}
        try:
            tg = datetime.fromisoformat(str(r.get("thoiGian")).replace("Z", "+00:00"))
        except ValueError:
            tg = timezone.now()
        LichSuGuiNhom.objects.create(
            nhom_id=str(b.get("nhomId") or "")[:64], ten_nhom=str(r.get("tenNhom") or "")[:255],
            ket_qua=str(r.get("ketQua") or "-")[:20], noi_dung=str(r.get("noiDung") or ""),
            so_anh=max(0, int(r.get("soAnh") or 0)), ghi_chu=str(r.get("ghiChu") or ""), thoi_gian=tg,
            tai_khoan_zalo=str(r.get("taiKhoanZalo") or "")[:255],
            nguoi_gui=NguoiDung.objects.filter(id=b.get("nguoiGuiId")).first(),
        )
        sv.tang_phien_ban([b.get("nguoiGuiId")])
    return JsonResponse({"ok": True})


# ---------- gửi vào nhóm Zalo ----------
@login_required
def trang_gui_nhom(request):
    """Trang gửi tin vào nhóm Zalo, dùng chung cho quản trị và nhân viên (gửi bằng Zalo của chính mình)."""
    u = request.user
    if not u.co_quyen("send"):
        raise PermissionDenied
    me = {"id": u.id, "name": str(u), "role": u.vai_tro, "roleLabel": u.get_vai_tro_display(), "staff": u.is_staff}
    return render(request, "guitin/gui_nhom.html", {"me": me})


@api("send", ["GET"])
def api_nhom_state(request):
    c = CauHinhNhom.cua(request.user)
    logs = list(nhat_ky_cho(request.user).filter(nguoi_dung=request.user).order_by("-id")[:200])
    return {**trang_thai_cho(request.user), "config": CaiDat.lay().to_dict(), "cauHinh": sv.cau_hinh_nhom_dict(c),
            "anhToiDa": sv.ANH_TOI_DA, "anhToiDaMb": sv.ANH_TOI_DA_MB,
            "logs": [sv.log_dict(n) for n in reversed(logs)], "version": sv.phien_ban(request.user)}


@api("send", ["GET"])
def api_nhom_danh_sach(request):
    """Các nhóm Zalo mà tài khoản Zalo của mình đang tham gia."""
    return worker.goi("GET", f"/groups?uid={request.user.id}", timeout=60)


@api("send", ["POST"])
def api_nhom_cau_hinh(request):
    """Lưu nhóm đích và nội dung gửi nhóm."""
    can_khong_dang_gui(request.user)
    c = CauHinhNhom.cua(request.user)
    b = request.json
    if "noiDung" in b:
        c.noi_dung = str(b.get("noiDung") or "").strip()[:5000]
    if "nhom" in b:
        nhom = b.get("nhom") or []
        if not isinstance(nhom, list):
            raise ValueError("Dữ liệu nhóm không hợp lệ")
        c.nhom = [{"id": str(g.get("id")), "ten": str(g.get("ten") or "")[:255]}
                  for g in nhom if isinstance(g, dict) and str(g.get("id") or "").isdigit()][:100]
    c.save()
    return {"cauHinh": sv.cau_hinh_nhom_dict(c)}


@api("send", ["POST"])
def api_nhom_anh(request):
    """Tải ảnh đính kèm lên (multipart, trường "anh", có thể nhiều file)."""
    can_khong_dang_gui(request.user)
    c = CauHinhNhom.cua(request.user)
    files = request.FILES.getlist("anh")
    if not files:
        raise ValueError("Chưa chọn ảnh")
    for f in files:
        sv.luu_anh(request.user, c, f)
    return {"cauHinh": sv.cau_hinh_nhom_dict(c)}


@api("send", ["POST"])
def api_nhom_anh_xoa(request):
    can_khong_dang_gui(request.user)
    c = CauHinhNhom.cua(request.user)
    sv.xoa_anh(request.user, c, str(request.json.get("ten") or ""))
    return {"cauHinh": sv.cau_hinh_nhom_dict(c)}


@login_required
def xem_anh_tin_nhan(request, pk):
    """Xem ảnh gửi kèm tin nhắn khách (quản trị cài; ai đăng nhập cũng xem được để biết sẽ gửi gì)."""
    from django.http import FileResponse, Http404
    from .models import AnhTinNhan
    a = AnhTinNhan.objects.filter(pk=pk).first()
    if not a or not a.tep or not a.tep.storage.exists(a.tep.name):
        raise Http404
    return FileResponse(a.tep.open("rb"))


@login_required
def xem_anh_nhom(request, ten):
    """Xem ảnh đính kèm của chính mình."""
    from django.http import FileResponse, Http404
    if ten not in CauHinhNhom.cua(request.user).anh:
        raise Http404
    f = sv.thu_muc_anh(request.user) / ten
    if not f.is_file():
        raise Http404
    return FileResponse(open(f, "rb"))


@api("send", ["POST"])
def api_nhom_gui(request):
    """Gửi nội dung + ảnh vào các nhóm đã chọn, bằng Zalo của chính người bấm."""
    u = request.user
    if worker.trang_thai(u.id).get("login") != "in":
        raise ValueError("Bạn chưa kết nối Zalo. Quét QR để kết nối Zalo của bạn trước khi gửi")
    c = CauHinhNhom.cua(u)
    if not c.nhom:
        raise ValueError("Chưa chọn nhóm nào để gửi")
    if not c.noi_dung and not c.anh:
        raise ValueError("Chưa có nội dung hay ảnh để gửi")
    worker.goi("POST", "/run-groups", {
        "uid": u.id, "groups": c.nhom, "noiDung": c.noi_dung,
        "anh": [f"anh-gui-nhom/{u.id}/{a}" for a in c.anh],
        "cfg": CaiDat.lay().to_dict(), "runBy": {"id": u.id, "name": str(u)},
    })
    return {"ok": True}


@api("send", ["GET"])
def api_nhom_lich_su(request):
    """100 lần gửi nhóm gần nhất (nhân viên chỉ thấy của mình)."""
    qs = LichSuGuiNhom.objects.select_related("nguoi_gui")
    if not request.user.co_quyen("all"):
        qs = qs.filter(nguoi_gui=request.user)
    return {"rows": [sv.lich_su_nhom_dict(h) for h in qs[:100]]}
