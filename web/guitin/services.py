"""Nghiệp vụ dùng chung cho trang gửi tin và trang quản trị."""
import itertools

from django.core.cache import cache
from django.db import transaction
from django.db.models import Count, Exists, Max, OuterRef, Q
from django.utils import timezone
from django.utils.timezone import localtime

from .models import CaiDat, KetQua, KhongLienHe, LichSuGui, NguoiDung, NguoiNhan, NhatKy, chuan_hoa_sdt, q_thanh_cong


# ---------- phiên bản danh sách: tăng mỗi khi danh sách / kết quả đổi để các trang đang mở tải lại ----------
# Phiên bản chung (đổi cài đặt, quản trị sửa danh sách: mọi trang tải lại) và phiên bản riêng từng người
# (kết quả gửi của 1 nhân viên chỉ làm trang của người đó và trang quản trị tải lại).
def _tang(key):
    try:
        cache.incr(key)
    except ValueError:
        cache.set(key, 1, None)


def tang_phien_ban(uids=None):
    if uids is None:
        _tang("ds_phien_ban")
        return
    for uid in {int(u) for u in uids if u}:
        _tang(f"ds_phien_ban:{uid}")
    _tang("ds_phien_ban:all")


def phien_ban(user=None):
    chung = cache.get_or_set("ds_phien_ban", 1, None)
    if user is None:
        return chung
    rieng = "all" if user.co_quyen("all") else int(user.id)
    return f"{chung}.{cache.get_or_set(f'ds_phien_ban:{rieng}', 1, None)}"


def ghi_log(noi_dung, muc="info", nguoi_dung=None):
    NhatKy.objects.create(noi_dung=noi_dung, muc=muc, nguoi_dung=nguoi_dung)
    don_nhat_ky()


NHAT_KY_GIU_NGAY = 60


def don_nhat_ky():
    """Mỗi ngày một lần: xoá nhật ký mức “Thông tin” cũ hơn NHAT_KY_GIU_NGAY ngày (cảnh báo / lỗi giữ lại)."""
    from datetime import timedelta
    if cache.add("don_nhat_ky", 1, 24 * 3600):
        NhatKy.objects.filter(muc="info", thoi_gian__lt=timezone.now() - timedelta(days=NHAT_KY_GIU_NGAY)).delete()


def log_dict(n):
    return {"id": n.id, "t": localtime(n.thoi_gian).strftime("%H:%M:%S %d/%m/%Y"), "text": n.noi_dung, "level": n.muc}


# ---------- phạm vi ----------
def pham_vi(user, chon="all"):
    """Trả về bộ lọc người phụ trách. Nhân viên luôn chỉ thấy số của mình.
    Quản trị / chỉ xem chọn được: all (tất cả), none (chưa giao), hoặc id nhân viên."""
    if not user.co_quyen("all"):
        return {"phu_trach": user}
    if chon in (None, "", "all"):
        return {}
    if chon == "none":
        return {"phu_trach__isnull": True}
    return {"phu_trach_id": int(chon)}


def nguoi_duoc_giao():
    """Những người có thể được giao số (nhân viên và quản trị đang hoạt động)."""
    return list(NguoiDung.objects.filter(is_active=True, vai_tro__in=[NguoiDung.NHAN_VIEN, NguoiDung.QUAN_TRI])
                .order_by("vai_tro", "ho_ten"))


def staff_dict(users):
    return [{"id": u.id, "name": u.ho_ten or u.email, "role": u.vai_tro} for u in users]


def ket_qua_dict(k):
    return {
        "ten": k.ten, "tenZalo": k.ten_zalo, "ketBan": k.ket_ban, "tinNhan": k.tin_nhan, "ghiChu": k.ghi_chu,
        "thoiGian": localtime(k.thoi_gian).strftime("%H:%M:%S %d/%m/%Y"),
        "nguoiGui": str(k.nguoi_gui) if k.nguoi_gui else "",
    }


def danh_sach(user, chon="all"):
    qs = NguoiNhan.objects.filter(**pham_vi(user, chon)).select_related("phu_trach")
    rows = list(qs)
    sdts = [r.sdt for r in rows if r.sdt]
    kq = {k.sdt: k for k in KetQua.objects.filter(sdt__in=sdts).select_related("nguoi_gui")}
    tu_choi = set(KhongLienHe.objects.filter(sdt__in=sdts).values_list("sdt", flat=True))
    return [{
        "id": r.id, "sdt": r.sdt_goc, "ten": r.ten, "noiDung": r.noi_dung, "duAn": r.du_an, "sdtChuan": r.sdt,
        "khoa": r.quan_tri_giao, "tuChoi": r.sdt in tu_choi,
        "phuTrach": r.phu_trach_id, "phuTrachTen": str(r.phu_trach) if r.phu_trach else None,
        "ketQua": ket_qua_dict(kq[r.sdt]) if r.sdt in kq else None,
    } for r in rows]


def cac_mau(text):
    """Tách nội dung thành các mẫu: mỗi mẫu cách nhau bằng một dòng chỉ có --- (giống worker)."""
    import re
    return [m.strip() for m in re.split(r"(?m)^\s*-{3,}\s*$", text or "") if m.strip()]


def xao(text):
    """Chọn ngẫu nhiên 1 mẫu, rồi thay {a|b|c} bằng 1 lựa chọn ngẫu nhiên. Worker làm y hệt khi gửi thật."""
    import random
    import re
    mau = cac_mau(text)
    t = random.choice(mau) if mau else ""
    for _ in range(10):
        t2 = re.sub(r"\{([^{}]*\|[^{}]*)\}", lambda m: random.choice(m.group(1).split("|")), t)
        if t2 == t:
            break
        t = t2
    return t


def dien(tpl, ten, du_an="", nhan_vien=""):
    """Thay {ten} bằng tên người nhận (trống thì thành "bạn"), {duan} bằng tên dự án, {nhanvien} bằng tên
    nhân viên phụ trách. Giống fill() ở worker và ở chung.js."""
    import re
    s = re.sub(r"\{ten\}", lambda m: ten or "bạn", tpl or "", flags=re.IGNORECASE)
    s = re.sub(r"\{duan\}", lambda m: du_an or "", s, flags=re.IGNORECASE)
    s = re.sub(r"\{nhanvien\}", lambda m: nhan_vien or "", s, flags=re.IGNORECASE)
    return s.strip()


def nhan_vien_cua(r):
    return str(r.phu_trach) if r.phu_trach_id and r.phu_trach else ""


def du_an_cua(r, cfg):
    """Tên dự án của 1 người nhận: dự án riêng của dòng, không có thì dự án mặc định trong cài đặt."""
    return (r.du_an or cfg.ten_du_an or "").strip()


def ten_goi_nho(tpl, ten, sdt, du_an=""):
    """Tên gợi nhớ đặt cho người được mời kết bạn (giống worker): {ten} {sdt} {duan}, thiếu phần nào thì bỏ phần đó."""
    import re
    s = re.sub(r"\{ten\}", lambda m: ten or "", tpl or "{ten} {sdt} {duan}", flags=re.IGNORECASE)
    s = re.sub(r"\{sdt\}", lambda m: sdt or "", s, flags=re.IGNORECASE)
    s = re.sub(r"\{duan\}", lambda m: du_an or "", s, flags=re.IGNORECASE)
    # Bỏ dấu nối thừa ở hai đầu khi thiếu tên, vd "{ten} - {sdt}" -> "0901234567"
    return " ".join(s.split()).strip(" -–|,.:")[:80] or sdt


def loi_moi_gui(r, cfg):
    """Lời mời kết bạn cho 1 người: của nhân viên phụ trách, trống thì của quản trị.
    Đã thay {duan}, {nhanvien}; còn {ten} để worker thay khi gửi."""
    cua_nv = r.phu_trach.loi_moi_ket_ban if r.phu_trach_id and r.phu_trach else ""
    return dien(cua_nv or cfg.loi_moi_ket_ban, "{ten}", du_an_cua(r, cfg), nhan_vien_cua(r))


def noi_dung_gui(r, cfg):
    """Nội dung tin nhắn sẽ gửi cho 1 người, theo thứ tự ưu tiên:
    nội dung riêng của dòng -> tin nhắn mặc định của nhân viên phụ trách -> nội dung mặc định của quản trị."""
    cua_nv = r.phu_trach.tin_nhan_mac_dinh if r.phu_trach_id and r.phu_trach else ""
    return dien(r.noi_dung or cua_nv or cfg.noi_dung_tin_nhan, r.ten, du_an_cua(r, cfg), nhan_vien_cua(r))


# ---------- hạn mức mỗi ngày của từng Zalo ----------
def dau_ngay():
    from datetime import datetime, time
    return timezone.make_aware(datetime.combine(timezone.localdate(), time.min))


def da_gui_hom_nay(user):
    """Số người Zalo của user đã xử lý hôm nay (mỗi lần gửi 1 dòng lịch sử, kể cả lỗi: Zalo vẫn tính lượt tìm số)."""
    return LichSuGui.objects.filter(nguoi_gui=user, thoi_gian__gte=dau_ngay()).count()


def han_muc_ngay(user, cfg):
    """Tối đa người/ngày của Zalo này: riêng của nhân viên, 0 thì theo cài đặt chung (0 = không giới hạn)."""
    return user.gioi_han_ngay or cfg.gioi_han_moi_ngay


def con_lai_hom_nay(user, cfg):
    """Hôm nay còn gửi được bao nhiêu người; None = không giới hạn."""
    toi_da = han_muc_ngay(user, cfg)
    return None if not toi_da else max(0, toi_da - da_gui_hom_nay(user))


def so_gui_toi_da(user, cfg):
    """Số người tối đa cho 1 lần bấm gửi: nhỏ hơn giữa “mỗi lần” và phần còn lại hôm nay. None = không giới hạn."""
    cac = [x for x in (cfg.gioi_han_moi_lan or None, con_lai_hom_nay(user, cfg)) if x is not None]
    return min(cac) if cac else None


def chua_gui_duoc(user, cfg):
    """Lý do Zalo của user chưa gửi được (đang tạm dừng vì lỗi, hết hạn mức hôm nay); None nếu gửi được."""
    if user.tam_dung_den and user.tam_dung_den > timezone.now():
        return (f"Zalo của bạn đang tạm dừng gửi đến {localtime(user.tam_dung_den):%H:%M %d/%m} vì đợt trước gặp nhiều lỗi "
                f"liên tiếp (có thể Zalo đang hạn chế tài khoản). Quản trị có thể bỏ tạm dừng trong Tài khoản.")
    if con_lai_hom_nay(user, cfg) == 0:
        return (f"Zalo của bạn đã gửi đủ {han_muc_ngay(user, cfg)} người hôm nay (hạn mức mỗi ngày để tránh bị Zalo khoá). "
                f"Mai gửi tiếp nhé.")
    return None


def cho_gui(user, chon, gioi_han=None):
    """Các số còn phải gửi trong phạm vi: SĐT hợp lệ, chưa có kết quả, không trong danh sách từ chối, không trùng,
    có nội dung (nếu gửi tin nhắn). Database tự loại các số đã gửi / từ chối."""
    cfg = CaiDat.lay()
    seen, out = set(), []
    qs = (NguoiNhan.objects.filter(**pham_vi(user, chon)).exclude(sdt__isnull=True)
          .exclude(Exists(KetQua.objects.filter(sdt=OuterRef("sdt"))))
          .exclude(Exists(KhongLienHe.objects.filter(sdt=OuterRef("sdt"))))
          .select_related("phu_trach"))
    for r in qs:
        if r.sdt in seen:
            continue
        if cfg.gui_tin_nhan and not noi_dung_gui(r, cfg):
            continue
        seen.add(r.sdt)
        out.append(r)
        if gioi_han and len(out) >= gioi_han:
            break
    return out


def _id(r):
    try:
        return int(r.get("id") or 0) or None
    except (TypeError, ValueError):
        return None


@transaction.atomic
def luu_danh_sach(user, rows):
    """Nhân viên lưu danh sách của mình, chỉ ghi phần thay đổi (giữ nguyên ngày thêm của từng số):
    - dòng có id: cập nhật; số quản trị giao thì không đổi được SĐT,
    - dòng mới: thêm, trừ số đã nằm trong danh sách của người khác hoặc khách đã từ chối liên hệ,
    - dòng bị bỏ: xoá, trừ số quản trị giao (nhân viên không xoá được).
    Trả về {"trungNguoiKhac": [...], "tuChoi": [...], "giuLai": số dòng quản trị giao không xoá}."""
    cua_toi = {r.id: r for r in NguoiNhan.objects.select_for_update().filter(phu_trach=user)}
    base = NguoiNhan.objects.aggregate(m=Max("thu_tu"))["m"] or 0
    cot = ("sdt_goc", "sdt", "ten", "noi_dung", "du_an")
    moi_sdt = {chuan_hoa_sdt(r.get("sdt")) for r in rows if _id(r) not in cua_toi} - {None}
    cua_nguoi_khac = set(NguoiNhan.objects.filter(sdt__in=moi_sdt).exclude(phu_trach=user).values_list("sdt", flat=True))
    tu_choi_set = set(KhongLienHe.objects.filter(sdt__in=moi_sdt).values_list("sdt", flat=True))
    giu, sua, moi, trung_khac, tu_choi = set(), [], [], [], []
    for r in rows:
        d = {"sdt_goc": str(r.get("sdt") or "").strip()[:32], "ten": str(r.get("ten") or "").strip()[:255],
             "noi_dung": str(r.get("noiDung") or "").strip()[:5000], "du_an": str(r.get("duAn") or "").strip()[:100]}
        d["sdt"] = chuan_hoa_sdt(d["sdt_goc"])
        cu = cua_toi.get(_id(r))
        if cu:
            giu.add(cu.id)
            if cu.quan_tri_giao:
                d["sdt_goc"], d["sdt"] = cu.sdt_goc, cu.sdt
            if any(getattr(cu, k) != d[k] for k in cot):
                for k in cot:
                    setattr(cu, k, d[k])
                sua.append(cu)
        elif d["sdt"] in cua_nguoi_khac:
            trung_khac.append(d["sdt"])
        elif d["sdt"] in tu_choi_set:
            tu_choi.append(d["sdt"])
        else:
            moi.append(NguoiNhan(phu_trach=user, thu_tu=base + len(moi) + 1, **d))
    bo = [r for rid, r in cua_toi.items() if rid not in giu]
    NguoiNhan.objects.filter(id__in=[r.id for r in bo if not r.quan_tri_giao]).delete()
    NguoiNhan.objects.bulk_update(sua, cot, batch_size=500)
    NguoiNhan.objects.bulk_create(moi, batch_size=500)
    tang_phien_ban([user.id])
    return {"trungNguoiKhac": trung_khac, "tuChoi": tu_choi, "giuLai": sum(1 for r in bo if r.quan_tri_giao)}


@transaction.atomic
def chia_deu():
    """Chia lần lượt các số chưa giao cho các nhân viên. Trả về (số đã chia, số nhân viên)."""
    nv = [u for u in nguoi_duoc_giao() if u.vai_tro == NguoiDung.NHAN_VIEN]
    if not nv:
        raise ValueError("Chưa có tài khoản nhân viên nào. Tạo trong trang quản trị → Tài khoản.")
    vong = itertools.cycle(nv)
    chua_giao = list(NguoiNhan.objects.filter(phu_trach__isnull=True))
    for r in chua_giao:
        r.phu_trach = next(vong)
        r.quan_tri_giao = True
    NguoiNhan.objects.bulk_update(chua_giao, ["phu_trach", "quan_tri_giao"], batch_size=500)
    tang_phien_ban()
    return len(chua_giao), len(nv)


# ---------- nhập danh sách ----------
def _bang_tu_file(data):
    """Đọc file thành danh sách dòng (mỗi dòng là list chuỗi). Hỗ trợ .xlsx, .xls (Excel 97-2003),
    .xls thực chất là bảng HTML (hay gặp khi xuất từ phần mềm khác) và CSV. File lỗi thì báo ValueError dễ hiểu."""
    import csv
    import io
    import re
    import unicodedata

    def o(v):
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)          # Excel lưu SĐT dạng số: 901234567.0 -> 901234567
        return str(v).strip()

    if data[:2] == b"PK":
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            return [[o(v) for v in r] for r in ws.iter_rows(values_only=True)]
        except Exception as e:  # noqa: BLE001
            raise ValueError("Không đọc được file Excel (.xlsx), file có thể bị lỗi. Mở file bằng Excel rồi lưu lại "
                             "(File → Save As → Excel Workbook .xlsx) và nhập lại.") from e

    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":          # Excel 97-2003 (.xls)
        try:
            import xlrd
            sh = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
            return [[o(v) for v in sh.row_values(i)] for i in range(sh.nrows)]
        except Exception as e:  # noqa: BLE001
            raise ValueError("Không đọc được file Excel đời cũ (.xls). Mở file bằng Excel rồi lưu thành .xlsx "
                             "(File → Save As → Excel Workbook) và nhập lại.") from e

    # Văn bản: thử UTF-8 trước, không được thì bảng mã tiếng Việt cũ của Windows
    for enc in ("utf-8-sig", "cp1258", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    text = unicodedata.normalize("NFC", text)    # cp1258 lưu dấu tách rời: ghép lại cho đúng chữ

    if "\x00" in text[:2000]:
        raise ValueError("File không phải Excel hay CSV. Hãy chọn file .xlsx, .xls hoặc .csv.")

    head = text.lstrip()[:2000].lower()
    if head.startswith("<") and ("<table" in text.lower()):    # .xls xuất từ web: thực chất là bảng HTML
        from html.parser import HTMLParser

        class Bang(HTMLParser):
            def __init__(self):
                super().__init__()
                self.rows, self.row, self.cell = [], None, None

            def handle_starttag(self, tag, attrs):
                if tag == "tr":
                    self.row = []
                elif tag in ("td", "th") and self.row is not None:
                    self.cell = []
                elif tag == "br" and self.cell is not None:
                    self.cell.append("\n")

            def handle_endtag(self, tag):
                if tag in ("td", "th") and self.row is not None and self.cell is not None:
                    self.row.append(re.sub(r"[ \t\r\f\v\xa0]+", " ", "".join(self.cell)).strip())
                    self.cell = None
                elif tag == "tr" and self.row is not None:
                    self.rows.append(self.row)
                    self.row = None

            def handle_data(self, d):
                if self.cell is not None:
                    self.cell.append(d)

        p = Bang()
        p.feed(text)
        return p.rows

    # CSV: chuẩn hoá xuống dòng (\r\n, \r kiểu Mac cũ) rồi đọc
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    first = text[:5000]
    delim = max((";", ",", "\t"), key=first.count)
    try:
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text, newline=""), delimiter=delim)]
    except csv.Error as e:
        raise ValueError(f"File CSV bị lỗi định dạng ({e}). Mở file bằng Excel rồi lưu thành .xlsx và nhập lại.") from e


def doc_file(data):
    """Đọc file Excel (.xlsx/.xls) hoặc CSV: cột SĐT / Số điện thoại / Phone, Tên / Họ tên / Name,
    và (không bắt buộc) Nội dung / Nội dung tin nhắn / Tin nhắn / Message.
    Dòng tiêu đề không cần nằm ở dòng đầu (bỏ qua dòng tên bảng phía trên)."""
    import unicodedata

    rows = _bang_tu_file(data)
    if not rows:
        return []

    def key(s):
        s = unicodedata.normalize("NFD", str(s).lower()).replace("đ", "d")
        return "".join(ch for ch in s if unicodedata.category(ch) != "Mn" and ch.isalnum())

    def la_sdt(h):
        return h in ("sdt", "sodienthoai", "dienthoai", "phone", "sodt", "didong", "mobile", "tel") or \
            h.startswith("sdt") or "dienthoai" in h or "didong" in h or "phone" in h

    def la_ten(h):
        return h in ("ten", "hoten", "hovaten", "name", "tenkhachhang", "khachhang", "fullname") or \
            (("ten" in h or "name" in h) and "zalo" not in h and "nhanvien" not in h and "phutrach" not in h)

    def la_nd(h):
        return h in ("noidung", "noidungtinnhan", "noidunggui", "tinnhan", "loinhan", "message", "content")

    def la_du_an(h):
        return h in ("duan", "tenduan", "project", "duanquantam", "sanpham")

    # Tìm dòng tiêu đề trong 10 dòng đầu
    for hi, r in enumerate(rows[:10]):
        head = [key(h) for h in r]
        sdt_cot = next((i for i, h in enumerate(head) if la_sdt(h)), None)
        if sdt_cot is not None:
            break
    else:
        raise ValueError("Không thấy cột SĐT. File cần có cột SĐT (hoặc Số điện thoại / Phone) và cột Tên. "
                         "Bấm “Tải file mẫu” để xem mẫu.")
    ten_cot = next((i for i, h in enumerate(head) if i != sdt_cot and la_ten(h)), None)
    nd_cot = next((i for i, h in enumerate(head) if la_nd(h)), None)
    da_cot = next((i for i, h in enumerate(head) if la_du_an(h)), None)
    out = []
    from .models import tach_sdt
    for r in rows[hi + 1:]:
        sdt = str(r[sdt_cot] if sdt_cot < len(r) else "").strip()
        if sdt.endswith(".0"):
            sdt = sdt[:-2]
        # Ô ghi nhiều số / số 11 chữ số cũ: lấy số hợp lệ đầu tiên
        cac_so = tach_sdt(sdt)
        if cac_so:
            sdt = cac_so[0]
        if sdt:
            out.append({"sdt": sdt, "ten": r[ten_cot] if ten_cot is not None and ten_cot < len(r) else "",
                        "noiDung": r[nd_cot] if nd_cot is not None and nd_cot < len(r) else "",
                        "duAn": r[da_cot] if da_cot is not None and da_cot < len(r) else ""})
    return out


def doc_van_ban(text):
    """Đọc danh sách dán vào: mỗi dòng 'SĐT, Tên, Nội dung' hoặc 2–4 cột copy từ Excel (cách nhau bằng Tab:
    SĐT, Tên, Nội dung, Dự án). Giống chức năng Dán ở trang nhân viên (chung.js)."""
    import re
    out = []
    for line in (text or "").splitlines():
        if not line.strip():
            continue
        if "\t" in line:
            parts = line.split("\t")
            if len(parts) > 4:      # nội dung có Tab: cột cuối là Dự án, phần giữa là nội dung
                parts = [parts[0], parts[1], "\t".join(parts[2:-1]), parts[-1]]
        else:
            parts = re.split(r"\s*[,;|]\s*", line.strip(), maxsplit=2)
        parts = [p.strip() for p in parts] + [""] * 4
        sdt = re.sub(r"[\s.\-]", "", parts[0])
        if re.fullmatch(r"\+?\d{8,12}", sdt):
            out.append({"sdt": sdt, "ten": parts[1], "noiDung": parts[2], "duAn": parts[3]})
    return out


@transaction.atomic
def nhap_them(rows, phu_trach=None, du_an=""):
    """Quản trị nhập thêm số vào cuối danh sách, bỏ qua số đã có (của bất kỳ ai) và khách đã từ chối liên hệ.
    Trả về {"them", "trung", "tuChoi", "daGui"} (daGui: trong số mới thêm, bao nhiêu số đã từng được gửi).
    du_an: tên dự án cho các dòng không có cột Dự án. Số được giao cho nhân viên thì nhân viên không xoá được."""
    da_co = set(NguoiNhan.objects.exclude(sdt__isnull=True).values_list("sdt", flat=True))
    tu_choi_set = set(KhongLienHe.objects.values_list("sdt", flat=True))
    base = NguoiNhan.objects.aggregate(m=Max("thu_tu"))["m"] or 0
    moi, trung, tu_choi = [], 0, 0
    for r in rows:
        sdt = chuan_hoa_sdt(r.get("sdt"))
        if sdt and sdt in da_co:
            trung += 1
            continue
        if sdt and sdt in tu_choi_set:
            tu_choi += 1
            continue
        if sdt:
            da_co.add(sdt)
        moi.append(NguoiNhan(thu_tu=base + len(moi) + 1, sdt_goc=str(r.get("sdt") or "").strip()[:32], sdt=sdt,
                             ten=str(r.get("ten") or "").strip()[:255], noi_dung=str(r.get("noiDung") or "").strip()[:5000],
                             du_an=(str(r.get("duAn") or "").strip() or du_an or "")[:100], phu_trach=phu_trach,
                             quan_tri_giao=phu_trach is not None))
    NguoiNhan.objects.bulk_create(moi, batch_size=500)
    da_gui = KetQua.objects.filter(sdt__in=[m.sdt for m in moi if m.sdt]).count()
    tang_phien_ban()
    return {"them": len(moi), "trung": trung, "tuChoi": tu_choi, "daGui": da_gui}


# ---------- thống kê cho bảng điều khiển ----------
def thong_ke():
    """Đếm theo nhân viên phụ trách; database tự đếm (không tải cả bảng lên).
    da_gui = gửi thành công (định nghĩa chung thanh_cong()); cho = còn phải gửi; loi = phần còn lại
    (gửi lỗi, số sai, khách từ chối liên hệ)."""
    co_kq = Exists(KetQua.objects.filter(sdt=OuterRef("sdt")))
    ok = Exists(KetQua.objects.filter(q_thanh_cong(), sdt=OuterRef("sdt")))
    tc = Exists(KhongLienHe.objects.filter(sdt=OuterRef("sdt")))
    dem = (NguoiNhan.objects.annotate(co_kq=co_kq, ok=ok, tc=tc).values("phu_trach_id")
           .annotate(tong=Count("id"), da_gui=Count("id", filter=Q(ok=True)),
                     cho=Count("id", filter=Q(sdt__isnull=False, co_kq=False, tc=False))).order_by())
    ten = {u.id: str(u) for u in NguoiDung.objects.filter(id__in=[d["phu_trach_id"] for d in dem if d["phu_trach_id"]])}
    nhom = {}
    tong = {"tong": 0, "cho": 0, "da_gui": 0, "loi": 0}
    for d in dem:
        key = d["phu_trach_id"] or 0
        g = nhom.setdefault(key, {"id": key, "ten": ten.get(key, "Chưa giao"), "tong": 0, "cho": 0, "da_gui": 0, "loi": 0})
        for x in (g, tong):
            x["tong"] += d["tong"]
            x["cho"] += d["cho"]
            x["da_gui"] += d["da_gui"]
            x["loi"] += d["tong"] - d["cho"] - d["da_gui"]
    # Nhân viên chưa được giao số nào vẫn hiện trong bảng
    for u in nguoi_duoc_giao():
        nhom.setdefault(u.id, {"id": u.id, "ten": str(u), "tong": 0, "cho": 0, "da_gui": 0, "loi": 0})
    rows = sorted(nhom.values(), key=lambda g: (g["id"] == 0, g["ten"]))
    return {"tong": tong, "nhanVien": rows}


# ---------- lịch sử gửi ----------
def loc_lich_su(user, q="", tt="", tu="", den="", nguoi=""):
    """Lọc lịch sử: nhân viên chỉ thấy lịch sử mình gửi. tu/den dạng YYYY-MM-DD."""
    from datetime import date, datetime, time, timedelta

    from django.utils.timezone import make_aware

    qs = LichSuGui.objects.select_related("nguoi_gui")
    if not user.co_quyen("all"):
        qs = qs.filter(nguoi_gui=user)
    elif nguoi:
        qs = qs.filter(nguoi_gui_id=int(nguoi))
    if q:
        qs = qs.filter(Q(sdt__icontains=q) | Q(ten__icontains=q) | Q(ten_zalo__icontains=q) | Q(noi_dung__icontains=q))
    if tt == "ok":
        qs = qs.filter(q_thanh_cong())
    elif tt == "loi":
        qs = qs.exclude(q_thanh_cong())
    # Lọc theo khoảng thời gian tính sẵn (không cần bảng múi giờ trong MySQL)
    try:
        if tu:
            qs = qs.filter(thoi_gian__gte=make_aware(datetime.combine(date.fromisoformat(tu), time.min)))
        if den:
            qs = qs.filter(thoi_gian__lt=make_aware(datetime.combine(date.fromisoformat(den) + timedelta(days=1), time.min)))
    except ValueError:
        raise ValueError("Ngày không hợp lệ") from None
    return qs


def lich_su_dict(h):
    return {
        "id": h.id, "sdt": h.sdt, "ten": h.ten, "tenZalo": h.ten_zalo, "ketBan": h.ket_ban, "tinNhan": h.tin_nhan,
        "noiDung": h.noi_dung, "ghiChu": h.ghi_chu, "thanhCong": h.thanh_cong, "taiKhoanZalo": h.tai_khoan_zalo,
        "thoiGian": localtime(h.thoi_gian).strftime("%H:%M:%S %d/%m/%Y"), "nguoiGui": str(h.nguoi_gui or ""),
    }


def o_excel(v):
    """Chống chèn công thức: ô chữ bắt đầu bằng = + - @ (hoặc Tab, xuống dòng) thêm dấu ' để Excel coi là chữ."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        return "'" + v
    return v


def dong_excel(ws, values):
    ws.append([o_excel(v) for v in values])


def excel_lich_su(qs):
    """File Excel lịch sử gửi (bytes)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "LichSuGui"
    ws.append(["Thời gian", "SĐT", "Họ tên", "Tên Zalo", "Kết bạn", "Tin nhắn", "Kết quả", "Nội dung đã gửi", "Ghi chú",
               "Nhân viên gửi", "Gửi từ Zalo"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FC4")
    for h in qs.order_by("thoi_gian", "id"):
        dong_excel(ws, [localtime(h.thoi_gian).strftime("%d/%m/%Y %H:%M:%S"), h.sdt, h.ten, h.ten_zalo, h.ket_ban, h.tin_nhan,
                   "Thành công" if h.thanh_cong else "Lỗi", h.noi_dung, h.ghi_chu, str(h.nguoi_gui or ""), h.tai_khoan_zalo])
    for col, w in zip("ABCDEFGHIJK", (20, 13, 22, 22, 11, 11, 12, 60, 40, 20, 22)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- gửi vào nhóm Zalo ----------
ANH_TOI_DA = 10
ANH_TOI_DA_MB = 5


def thu_muc_anh(user):
    """Ảnh gửi nhóm của từng tài khoản: data/anh-gui-nhom/<id>/ (worker đọc cùng thư mục này)."""
    from django.conf import settings
    return settings.ROOT_DIR / "data" / "anh-gui-nhom" / str(int(user.id))


def loai_anh(data):
    """Đuôi file theo nội dung ảnh (jpg/png/webp), None nếu không phải ảnh hỗ trợ."""
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if data[:3] == b"\xff\xd8\xff":
        return "jpg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    return None


def luu_anh(user, cau_hinh, f):
    """Lưu 1 ảnh tải lên vào thư mục của người dùng và thêm vào cấu hình gửi nhóm."""
    import uuid
    if len(cau_hinh.anh) >= ANH_TOI_DA:
        raise ValueError(f"Tối đa {ANH_TOI_DA} ảnh mỗi lần gửi")
    if f.size > ANH_TOI_DA_MB * 1024 * 1024:
        raise ValueError(f"Ảnh “{f.name}” lớn hơn {ANH_TOI_DA_MB} MB")
    data = f.read()
    ext = loai_anh(data)
    if not ext:
        raise ValueError(f"“{f.name}” không phải ảnh JPG, PNG hoặc WEBP")
    thu_muc = thu_muc_anh(user)
    thu_muc.mkdir(parents=True, exist_ok=True)
    ten = f"{uuid.uuid4().hex}.{ext}"
    (thu_muc / ten).write_bytes(data)
    cau_hinh.anh = [*cau_hinh.anh, ten]
    cau_hinh.save(update_fields=["anh", "cap_nhat"])


def xoa_anh(user, cau_hinh, ten):
    if ten not in cau_hinh.anh:
        raise ValueError("Không tìm thấy ảnh")
    cau_hinh.anh = [a for a in cau_hinh.anh if a != ten]
    cau_hinh.save(update_fields=["anh", "cap_nhat"])
    (thu_muc_anh(user) / ten).unlink(missing_ok=True)


# ---------- ảnh gửi kèm tin nhắn của từng nhân viên ----------
def thu_muc_anh_nv(user):
    """data/anh-tin-nhan/nv-<id>/ (nằm trong thư mục ảnh worker được phép đọc)."""
    from django.conf import settings
    return settings.ROOT_DIR / "data" / "anh-tin-nhan" / f"nv-{int(user.id)}"


def luu_anh_nv(user, f):
    import uuid
    if len(user.anh_tin_nhan) >= ANH_TOI_DA:
        raise ValueError(f"Tối đa {ANH_TOI_DA} ảnh")
    if f.size > ANH_TOI_DA_MB * 1024 * 1024:
        raise ValueError(f"Ảnh “{f.name}” lớn hơn {ANH_TOI_DA_MB} MB")
    data = f.read()
    ext = loai_anh(data)
    if not ext:
        raise ValueError(f"“{f.name}” không phải ảnh JPG, PNG hoặc WEBP")
    thu_muc = thu_muc_anh_nv(user)
    thu_muc.mkdir(parents=True, exist_ok=True)
    ten = f"{uuid.uuid4().hex}.{ext}"
    (thu_muc / ten).write_bytes(data)
    user.anh_tin_nhan = [*user.anh_tin_nhan, ten]
    user.save(update_fields=["anh_tin_nhan"])


def xoa_anh_nv(user, ten):
    if ten not in user.anh_tin_nhan:
        raise ValueError("Không tìm thấy ảnh")
    user.anh_tin_nhan = [a for a in user.anh_tin_nhan if a != ten]
    user.save(update_fields=["anh_tin_nhan"])
    (thu_muc_anh_nv(user) / ten).unlink(missing_ok=True)


def anh_gui_kem(user, cfg):
    """Ảnh gửi kèm tin nhắn khi user bấm gửi: ảnh riêng của user, chưa có thì ảnh của quản trị (đường dẫn trong data/)."""
    if user.anh_tin_nhan:
        return [f"anh-tin-nhan/nv-{int(user.id)}/{a}" for a in user.anh_tin_nhan]
    return cfg.to_dict()["anh"]


def cau_hinh_nhom_dict(c):
    return {"noiDung": c.noi_dung, "nhom": c.nhom, "anh": c.anh}


def lich_su_nhom_dict(h):
    return {
        "id": h.id, "nhomId": h.nhom_id, "tenNhom": h.ten_nhom, "ketQua": h.ket_qua, "thanhCong": h.thanh_cong,
        "noiDung": h.noi_dung, "soAnh": h.so_anh, "ghiChu": h.ghi_chu, "taiKhoanZalo": h.tai_khoan_zalo,
        "thoiGian": localtime(h.thoi_gian).strftime("%H:%M:%S %d/%m/%Y"), "nguoiGui": str(h.nguoi_gui or ""),
    }
